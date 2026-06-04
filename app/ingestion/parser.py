import io
from dataclasses import dataclass


@dataclass
class ParsedDocument:
    pages: list[str]  # one entry per page; plain text for TXT has a single entry


def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return ParsedDocument(pages=pages)


def parse_docx(data: bytes) -> ParsedDocument:
    import docx

    doc = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    return ParsedDocument(pages=[text])


def parse_txt(data: bytes) -> ParsedDocument:
    return ParsedDocument(pages=[data.decode("utf-8", errors="replace")])


def parse(data: bytes, filename: str) -> ParsedDocument:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return parse_pdf(data)
    elif ext == "docx":
        return parse_docx(data)
    else:
        return parse_txt(data)
